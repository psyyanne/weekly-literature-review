"""Validate the compact workbook, deep-reading DOCX files, IDs and active scope."""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from common import LITERATURE_COLUMNS, SHEET_SCHEMAS, project_root, worksheet_records

ID_RE = re.compile(r"^\d{4}-W\d{2}-[CFXO]\d+$")


def validate(root: Path) -> dict[str, object]:
    errors: list[str] = []
    workbook = root / "library/master_literature_review.xlsx"
    wb = load_workbook(workbook, read_only=True, data_only=True)
    if wb.sheetnames != list(SHEET_SCHEMAS):
        errors.append(f"工作簿必须只有四张表，实际为: {wb.sheetnames}")
    for sheet_name, columns in SHEET_SCHEMAS.items():
        if sheet_name not in wb.sheetnames or [cell.value for cell in wb[sheet_name][1]] != columns:
            errors.append(f"{sheet_name}表头与技能定义不一致")
    rows = {name: worksheet_records(wb[name]) for name in wb.sheetnames}
    wb.close()
    literature = rows.get("文献表", [])
    ids = [str(r.get("文献编号") or "") for r in literature]
    if len(ids) != len(set(ids)):
        errors.append("文献编号重复")
    if any(not ID_RE.fullmatch(value) for value in ids):
        errors.append("存在不符合YYYY-Www-C/F/X/O编号规则的文献编号")
    dois = [str(r.get("DOI") or "").lower() for r in literature if str(r.get("DOI") or "") not in {"", "待核验"}]
    if len(dois) != len(set(dois)):
        errors.append("DOI重复")
    literature_by_id = {str(r.get("文献编号")): r for r in literature}
    deep_docs = {p.name: p for p in (root / "outputs/deep_reading").glob("*.docx")} if (root / "outputs/deep_reading").exists() else {}
    for row in literature:
        path = str(row.get("精读文档路径") or "")
        if row.get("精读状态") == "已完成全文精读":
            target = root / path if not Path(path).is_absolute() else Path(path)
            if not target.is_file():
                errors.append(f"{row.get('文献编号')}精读DOCX不存在")
            else:
                text = "\n".join(p.text for p in Document(target).paragraphs)
                if str(row.get("文献编号")) not in text:
                    errors.append(f"{row.get('文献编号')}未在DOCX正文出现")
                count = len(re.findall(re.escape(str(row.get("文献编号"))) + r"-T\d{2}", text))
                if not 3 <= count <= 5:
                    errors.append(f"{row.get('文献编号')}DOCX备选题目数量为{count}")
    topic_rows = rows.get("选题池", [])
    for topic in topic_rows:
        source = str(topic.get("来源文献编号") or "")
        if source not in literature_by_id:
            errors.append(f"{topic.get('选题编号')}来源文献不存在")
        if not str(topic.get("选题编号") or "").startswith(source + "-T"):
            errors.append(f"{topic.get('选题编号')}未按来源文献编号生成")
    known_ids = set(literature_by_id)
    for trend in rows.get("研究趋势", []):
        support = [x.strip() for x in re.split(r"[；;,，\s]+", str(trend.get("支持文献编号") or "")) if x.strip()]
        if not support or any(x not in known_ids for x in support):
            errors.append(f"{trend.get('趋势编号')}缺少有效支持文献编号")
    active = [root / ".agents/skills/weekly-literature-review", root / "config", root / "README.md"]
    legacy_tokens = ["J" + "CR", "JI" + "F", "\u5f71\u54cd\u56e0\u5b50", "\u535a\u58eb\u8bba\u6587", "\u7ec4\u4f1a\u5468\u62a5", "\u7ec4\u4f1a\u6c47\u62a5", "P" + "PT", "\u53e3\u64ad\u7a3f", "\u8001\u5e08\u8ffd\u95ee"]
    forbidden = re.compile("|".join(re.escape(x) for x in legacy_tokens), re.I)
    hits = []
    for base in active:
        files = [base] if base.is_file() else [p for p in base.rglob("*") if p.is_file() and p.suffix.lower() in {".md", ".yaml", ".yml", ".py", ".json"}]
        for file in files:
            text = file.read_text(encoding="utf-8", errors="ignore")
            if forbidden.search(text):
                hits.append(str(file.relative_to(root)))
    if hits:
        errors.append("活动技能/配置中仍有已删除功能或字段: " + ", ".join(hits))
    expected_scripts = {"common.py", "search_and_verify_literature.py", "update_literature_workbook.py", "build_deep_read_docx.py", "update_topic_pool.py", "update_research_trends.py", "validate_outputs.py"}
    script_dir = root / ".agents/skills/weekly-literature-review/scripts"
    extras = sorted(p.name for p in script_dir.glob("*.py") if p.name not in expected_scripts)
    if extras:
        errors.append("非当前流程脚本仍位于活动路径: " + ", ".join(extras))
    return {"valid": not errors, "errors": errors, "counts": {"literature": len(literature), "deep_docs": len(deep_docs), "topics": len(topic_rows), "trends": len(rows.get("研究趋势", []))}}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project-root")
    args = parser.parse_args()
    result = validate(project_root(args.project_root))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["valid"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
