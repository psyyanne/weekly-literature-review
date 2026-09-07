"""Build one evidence-bounded DOCX deep-reading record per full-text paper."""
from __future__ import annotations

import argparse
import json
import re
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import load_workbook

from common import atomic_save, project_root, worksheet_records


TOPIC_SEEDS: dict[str, list[dict[str, str]]] = {
    "mental health chatbots": [
        {"title": "语境错配如何通过信任更新影响心理健康聊天机器人的持续使用", "gap": "现有证据多为短期体验，缺少对语境错配、信任变化和留存的纵向机制检验。", "theory": "信任校准与治疗联盟", "vars": "语境错配、感知准确性、信任更新、持续使用", "sample": "有真实聊天机器人使用经历的成年人", "method": "两阶段纵向体验抽样与多层模型", "advance": "把设计特征与时间序列中的信任变化连接起来", "novelty": "区分准确性、共情和危机安全三个信号", "feasibility": "可用公开系统或模拟对话材料实施", "verify": "需要预先定义语境错配编码与危机情境伦理边界"},
        {"title": "人类升级提示能否阻断心理健康聊天机器人的错误信任", "gap": "关于人类转介、透明纠错和持续参与的证据尚未形成可比较的实验设计。", "theory": "自动化信任与人类监督", "vars": "升级提示、透明度、信任、建议采纳、求助意愿", "sample": "高压力或低心理健康素养成人", "method": "情境实验加建议采纳任务", "advance": "从静态满意度转向安全决策和人类升级行为", "novelty": "比较无提示、透明提示和强制转介三种安全策略", "feasibility": "可通过标准化对话片段和在线实验完成", "verify": "需要核验材料是否诱发真实心理风险"},
        {"title": "生成式心理健康系统的共情表达何时促进而非削弱恰当求助", "gap": "共情知觉可能提升联盟，也可能掩盖系统能力边界。", "theory": "治疗联盟与能力—温暖模型", "vars": "共情表达、能力知觉、温暖知觉、求助意愿、能力校准", "sample": "不同AI素养水平的心理健康服务使用者", "method": "2×2实验与调节中介模型", "advance": "检验共情和能力边界的交互而非单独主效应", "novelty": "把恰当求助作为结果而非单纯使用意愿", "feasibility": "变量和操纵可以从原综述证据直接转化", "verify": "需要确认共情材料的文化适配性"},
    ],
    "ai models explain learning": [
        {"title": "错误签名能否区分AI认知模型的拟合与真正心理解释", "gap": "模型能复现平均表现不等于能解释错误模式和学习轨迹。", "theory": "认知模型效度与错误签名", "vars": "错误类型、学习轨迹、模型参数、解释效度", "sample": "多阶段学习任务中的人类参与者", "method": "预注册行为实验与模型比较", "advance": "将错误签名作为模型外部效度的核心检验", "novelty": "同时比较平均反应、个体差异和错误序列", "feasibility": "可使用公开认知任务和可复现模型", "verify": "需要预先定义模型失败而非事后挑选错误"},
        {"title": "人类学习轨迹与AI模型参数的跨任务可迁移性", "gap": "单任务拟合无法回答模型参数是否代表稳定心理机制。", "theory": "潜在过程与计算认知建模", "vars": "学习率、探索率、认知负荷、跨任务迁移", "sample": "完成两种结构相近学习任务的成人", "method": "跨任务层级贝叶斯模型与预测检验", "advance": "从描述性拟合推进到跨任务预测", "novelty": "区分参数稳定性和任务特定策略", "feasibility": "任务与模型可基于原文效度标准扩展", "verify": "需要样本量支持参数不确定性估计"},
        {"title": "把生态效度纳入AI认知模型解释资格的联合标准", "gap": "理论基础、构念对应和生态效度往往被分开报告。", "theory": "构念效度与生态效度", "vars": "理论对齐、过程机制、错误签名、生态效度", "sample": "实验室任务与自然学习场景的配对样本", "method": "多标准效度评分与外部验证", "advance": "形成可操作的模型评估清单", "novelty": "把课堂/真实环境预测纳入认知模型评价", "feasibility": "可先做方法学验证再开展大样本研究", "verify": "需要统一不同任务中的效度指标"},
    ],
    "observer perceptions": [
        {"title": "生成式AI使用标签如何改变温暖、能力与人际信任的相对权重", "gap": "现有研究发现信任变化，但社会信号之间的权重动态仍不清楚。", "theory": "社会判断的温暖—能力模型", "vars": "AI使用标签、温暖、能力、人际信任、AI素养", "sample": "观察他人协作或沟通的成人样本", "method": "标签实验与潜变量路径模型", "advance": "比较标签对社会判断维度的差异化影响", "novelty": "将观察者AI素养作为校准条件", "feasibility": "可直接复用原文信任和社会判断量表", "verify": "需要检验标签是否引发责任归因变化"},
        {"title": "AI使用透明度何时提升而非降低合作信任", "gap": "透明披露可能提升诚信判断，也可能触发能力怀疑。", "theory": "信任校准与归因理论", "vars": "披露透明度、能力、诚信、合作信任、任务难度", "sample": "人机团队决策参与者", "method": "2×2任务难度实验与条件过程分析", "advance": "将任务难度作为AI标签效果的边界条件", "novelty": "区分对AI系统的信任和对人类伙伴的信任", "feasibility": "可用短时人机协作任务", "verify": "需要区分披露内容与披露时机"},
        {"title": "AI素养是否使观察者更能校准对生成式AI使用者的信任", "gap": "AI素养可能提高辨别能力，也可能放大刻板印象。", "theory": "社会信息加工与元认知校准", "vars": "AI素养、使用标签、证据质量、信任校准", "sample": "不同AI素养水平的观察者", "method": "多轮证据更新任务与校准曲线", "advance": "用校准误差而非单一信任评分评估AI素养作用", "novelty": "把社会判断与客观证据识别结合", "feasibility": "可设计标准化行为证据序列", "verify": "需要建立客观正确性基准"},
    ],
    "six-facet artificial intelligence literacy": [
        {"title": "六面AI素养结构能否预测真实错误识别而非自我评价", "gap": "量表高分是否对应客观判断表现仍缺乏行为验证。", "theory": "AI素养与元认知校准", "vars": "六面AI素养、错误识别、信心、校准误差", "sample": "青少年、青年与中年成人", "method": "量表加行为判断任务与多组结构方程", "advance": "将主观素养与客观能力同时建模", "novelty": "比较不同年龄段的素养—校准关系", "feasibility": "量表已有题目和维度，行为任务可独立开发", "verify": "需要确认题目版权和不同年龄段测量等值性"},
        {"title": "AI素养干预能否减少过度依赖并保留有效使用", "gap": "提高素养不一定自动带来恰当依赖，干预后的行为结果尚不明确。", "theory": "恰当依赖与自我调节学习", "vars": "AI素养干预、依赖校准、错误识别、任务绩效", "sample": "需要使用生成式AI完成学习任务的学生", "method": "随机干预与延迟后测", "advance": "从量表开发推进到可验证的行为干预", "novelty": "同时考察过度依赖与依赖不足", "feasibility": "可将量表维度转为短时教学模块", "verify": "需要确定干预剂量和迁移任务"},
        {"title": "AI素养六个维度在高风险决策中的相对贡献", "gap": "总分掩盖了不同素养维度对安全决策的差异。", "theory": "多维能力与风险决策", "vars": "六维素养、风险识别、建议采纳、责任归因", "sample": "医疗、教育或公共服务决策者", "method": "多组路径模型与相对权重分析", "advance": "识别面向情境的最小素养组合", "novelty": "以高风险决策而非一般态度为结果", "feasibility": "可在不提供真实医疗建议的模拟任务中实施", "verify": "需要场景专家评审和伦理审查"},
    ],
    "cognitive capability and behavioral exposure": [
        {"title": "认知能力与AI行为暴露何时导致促进而非学术诚信风险", "gap": "双路径模型的复制结果不稳定，边界条件需要预注册验证。", "theory": "能力—暴露双路径与自我调节", "vars": "认知能力、行为暴露、感知促进、诚信风险", "sample": "不同学业阶段学生", "method": "预注册多样本复现与等值性检验", "advance": "把路径不稳定性作为理论边界而非噪声", "novelty": "同时解释效率收益和诚信风险", "feasibility": "可复用原文测量并增加客观行为指标", "verify": "需要改进模型拟合和样本代表性"},
        {"title": "客观AI能力能否缓冲生成式AI暴露对诚信风险的影响", "gap": "感知能力和客观能力可能对风险路径产生不同方向的调节。", "theory": "元认知校准与自我控制", "vars": "客观能力、感知能力、AI暴露、诚信风险", "sample": "大学生学习者", "method": "能力测验、使用日志与纵向追踪", "advance": "分离感知能力与客观能力的作用", "novelty": "以使用过程而非单次自报暴露测量行为", "feasibility": "可采用匿名日志和标准化任务", "verify": "需要解决隐私、日志同意和共同方法偏差"},
        {"title": "生成式AI使用中的保护性认知策略如何形成", "gap": "现有模型强调促进和风险结果，较少解释用户如何主动校准使用。", "theory": "自我调节与保护性动机", "vars": "错误检查、来源监控、提示策略、诚信行为", "sample": "长期使用生成式AI的学习者", "method": "经验取样加策略训练实验", "advance": "从结果模型推进到可干预的过程机制", "novelty": "把来源监控和错误检查作为中间过程", "feasibility": "可基于现有双路径模型开发短训", "verify": "需要客观记录策略执行而非只测态度"},
    ],
}


def topic_candidates(deep: dict[str, Any]) -> list[dict[str, str]]:
    title = str(deep.get("英文标题") or "").lower()
    chosen = next((items for key, items in TOPIC_SEEDS.items() if key in title), None)
    if chosen is None:
        base = str(deep.get("研究问题") or deep.get("英文标题") or "AI心理学研究")
        theory = str(deep.get("理论名称与中文含义") or "待核验")
        variables = str(deep.get("核心概念及作者定义") or "核心心理变量")
        chosen = [
            {"title": f"从{base}检验理论边界", "gap": "原文研究结果需要在不同样本或情境中复核。", "theory": theory, "vars": variables, "sample": "与原文不同的心理学样本", "method": "预注册复现与稳健性分析", "advance": "扩展原文样本或设计边界", "novelty": "把原文局限转为可检验问题", "feasibility": "中等，需先完成材料核验", "verify": "需核验原文材料、量表和样本量"},
            {"title": f"{base}中的用户校准机制", "gap": "原文通常聚焦系统或变量关联，较少直接测量用户是否形成恰当依赖。", "theory": "信任校准、元认知与自我调节", "vars": "信任、信心、错误识别、建议采纳、任务表现", "sample": "具有AI工具使用经验的学生或成人用户", "method": "行为任务结合信心评分和校准曲线", "advance": "把态度或评价推进到可观察的校准行为", "novelty": "同时区分过度依赖、依赖不足和恰当依赖", "feasibility": "可用在线任务和标准化AI输出材料实施", "verify": "需要建立客观正确性基准和风险等级"},
            {"title": f"{base}的纵向变化和干预可塑性", "gap": "横断或评论证据不能说明机制是否随经验变化或可被训练改变。", "theory": "学习迁移、技术适应与保护性动机", "vars": "AI经验、策略使用、风险识别、心理结果、持续使用", "sample": "连续使用生成式AI工具的学习者或助人服务使用者", "method": "纵向追踪或短期干预加延迟后测", "advance": "从一次性判断推进到时间变化和干预效果", "novelty": "把AI素养、信任和安全行为放在同一动态模型中", "feasibility": "可采用两到三波在线追踪和日志自报", "verify": "需要处理隐私同意、样本流失和模型版本变化"},
        ]
    return chosen[:5]


def setup_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "SimSun")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "SimSun")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, color in (("Title", 22, "0B2545"), ("Heading 1", 15, "2E74B5"), ("Heading 2", 12, "1F4D78")):
        style = doc.styles[name]
        style.font.name = "SimSun"
        style._element.rPr.rFonts.set(qn("w:ascii"), "SimSun")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "SimSun")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True


def usable(value: Any) -> bool:
    text = str(value or "").strip()
    return bool(text) and "未报告" not in text and "需全文核验" not in text and "全文核验" not in text


def add_field(doc: Document, label: str, value: Any) -> bool:
    if not usable(value):
        return False
    p = doc.add_paragraph()
    p.add_run(f"{label}：").bold = True
    p.add_run(str(value).strip())
    return True


def slug(text: str) -> str:
    text = re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")
    return (text[:70] or "paper")


def build_one(root: Path, record: dict[str, Any], doc_id: str, topics: list[dict[str, str]]) -> Path:
    out_dir = root / "outputs/deep_reading"
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"{doc_id}_{slug(str(record.get('英文标题') or 'paper'))}.docx"
    doc = Document()
    setup_document(doc)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run(f"{doc_id}｜{record.get('英文标题') or doc_id}")
    doc.add_paragraph("基于实际全文的证据边界精读记录。作者观点、证据解释与AI延伸分析分开记录；仅保留论文实际呈现的信息。")
    doc.add_heading("一、基本信息", level=1)
    for label, key in (("文献编号", None), ("文章类型", "文章类型"), ("英文题名", "英文标题"), ("中文辅助译题", "中文标题"), ("完整作者", "作者"), ("第一位作者", "第一作者"), ("期刊", "期刊"), ("DOI", "DOI"), ("正式发表日期", "发表日期"), ("APA格式参考文献", "APA参考文献"), ("全文访问来源", "全文来源"), ("精读日期", "阅读日期")):
        add_field(doc, label, doc_id if label == "文献编号" else record.get(key))
    sections = [
        ("二、一句话概括", [("研究内容", "研究问题"), ("方法", "统计方法详情"), ("主要结论", "作者主要结论")]),
        ("三、研究背景与研究缺口", [("领域位置", "研究背景与领域位置"), ("重要性", "研究问题的重要性"), ("现有研究与缺口", "现有研究与缺口"), ("作者动机", "作者开展研究的原因")]),
        ("四、理论基础", [("理论名称与含义", "理论名称与中文含义"), ("核心观点", "理论核心观点"), ("本文作用", "理论在本文中的作用"), ("与问题/假设连接", "理论与问题及假设的连接"), ("可进一步检验", "可进一步检验的问题")]),
        ("五、核心概念", [("作者定义", "核心概念及作者定义"), ("操作性定义", "核心概念操作性定义"), ("相近概念区分", "相近概念区分"), ("模型位置", "概念在模型中的位置")]),
        ("六、研究问题和假设", [("研究问题", "研究问题"), ("假设及依据", "研究假设及理论依据"), ("支持情况", "假设支持情况"), ("探索性分析", "探索性分析"), ("不显著假设", "不显著假设")]),
        ("七、研究模型", [("变量关系", "变量理论关系"), ("自变量", "自变量"), ("因变量", "因变量"), ("中介/调节/控制", "中介变量"), ("操作化", "变量操作化"), ("可重绘模型", "可重绘研究模型")]),
        ("八、样本和研究设计", [("样本与设计", "分研究样本与设计"), ("招募与纳排", "招募与纳排"), ("随机与条件", "随机分配与实验条件"), ("预注册与开放科学", "预注册与开放科学")]),
        ("九、量表和测量", [("量表/测量详情", "量表与测量工具详情")]),
        ("十、实验操纵", [("操纵详情", "实验操纵详情")]),
        ("十一、统计方法", [("统计方法详情", "统计方法详情"), ("数据预处理", "数据预处理"), ("缺失与异常", "缺失值与异常值处理"), ("拟合与稳健性", "模型拟合与稳健性"), ("因果边界", "因果推断边界")]),
        ("十二、具体结果", [("具体统计结果", "具体统计结果"), ("不显著结果", "不显著结果"), ("效应量与区间", "效应量与置信区间"), ("中介与调节", "中介与调节结果"), ("组间差异", "组间差异"), ("证据位置", "页码表格图形证据位置")]),
        ("十三、讨论和贡献", [("作者解释", "作者对结果的解释"), ("理论贡献", "理论贡献"), ("方法贡献", "方法贡献"), ("实践意义", "实践意义"), ("作者未来方向", "作者未来方向")]),
        ("十四、局限和延伸分析", [("作者明确局限", "作者明确局限"), ("合理解释", "基于证据的合理解释"), ("AI延伸分析", "AI延伸分析"), ("替代解释", "替代解释"), ("样本/测量/设计/统计/推广限制", "样本限制")]),
        ("十五、可复用研究资料", [("理论", "理论资料价值"), ("变量", "变量资料价值"), ("测量", "测量资料价值"), ("实验设计", "实验设计资料价值"), ("样本", "样本资料价值"), ("统计方法", "统计方法资料价值"), ("可复制材料", "可复制材料"), ("谨慎采用", "需要谨慎采用的内容")]),
    ]
    for heading, fields in sections:
        values = [(label, record.get(key)) for label, key in fields if usable(record.get(key))]
        if not values:
            continue
        doc.add_heading(heading, level=1)
        for label, value in values:
            add_field(doc, label, value)
    doc.add_heading("十六、基于单篇文献的备选研究题目", level=1)
    for i, topic in enumerate(topics, 1):
        doc.add_heading(f"{doc_id}-T{i:02d}｜{topic['title']}", level=2)
        for label in ("gap", "theory", "vars", "sample", "method", "advance", "novelty", "feasibility", "verify"):
            add_field(doc, {"gap":"研究缺口", "theory":"理论基础", "vars":"核心变量", "sample":"研究对象", "method":"建议研究方法", "advance":"相比原文的推进", "novelty":"主要创新点", "feasibility":"可行性", "verify":"需要继续核验的问题"}[label], topic[label])
    doc.save(path)
    return path


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("bundle_json")
    parser.add_argument("--project-root")
    parser.add_argument("--doc-id")
    args = parser.parse_args()
    root = project_root(args.project_root)
    bundle = json.loads(Path(args.bundle_json).read_text(encoding="utf-8"))
    wb = load_workbook(root / "library/master_literature_review.xlsx")
    lit_rows = worksheet_records(wb["文献表"])
    by_doi = {str(r.get("DOI") or "").lower(): r for r in lit_rows}
    deep_rows = bundle.get("deep_readings") or []
    written = []
    for deep in deep_rows:
        row = by_doi.get(str(deep.get("DOI") or "").lower())
        if not row:
            continue
        doc_id = str(row["文献编号"])
        if args.doc_id and doc_id != args.doc_id:
            continue
        path = build_one(root, deep, doc_id, topic_candidates(deep))
        for excel_row in range(2, wb["文献表"].max_row + 1):
            if wb["文献表"].cell(excel_row, 1).value == doc_id:
                wb["文献表"].cell(excel_row, 13).value = "已完成全文精读"
                wb["文献表"].cell(excel_row, 14).value = str(path.relative_to(root)).replace("\\", "/")
        written.append(str(path))
    backup = atomic_save(wb, root / "library/master_literature_review.xlsx", make_backup=True)
    print(json.dumps({"written": written, "count": len(written), "backup": str(backup) if backup else None}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
