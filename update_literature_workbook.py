"""Migrate/upsert literature records into the four-sheet cumulative workbook."""
from __future__ import annotations

import argparse
import json
import re
import shutil
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from common import (
    LITERATURE_COLUMNS,
    SHEET_SCHEMAS,
    append_change,
    atomic_save,
    create_compact_workbook,
    project_root,
    record_key,
    worksheet_records,
)


def category_for(record: dict[str, Any]) -> str:
    journal = str(record.get("期刊") or record.get("来源或期刊") or "").lower()
    kind = str(record.get("文章类型") or record.get("材料类型") or "").lower()
    if "frontiers" in journal:
        return "F"
    if any(token in kind for token in ("预印本", "preprint", "会议", "conference")) or not str(record.get("DOI") or "").strip() or str(record.get("是否正式发表") or "是") in {"否", "未正式发表"}:
        return "O"
    cross_journals = ("digital medicine", "social sciences", "medical", "medicine", "health", "jmir", "jama network open", "education technology", "computer", "information system", "management", "human-computer", "hci")
    if any(token in journal for token in cross_journals):
        return "X"
    return "C"


def read_old_records(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    for name in wb.sheetnames:
        if name in {"选题池", "研究趋势", "修改日志"}:
            continue
        for row in worksheet_records(wb[name]):
            if any(key in row for key in ("文献ID", "文献编号", "DOI", "英文标题", "英文题名")):
                records.append(row)
    wb.close()
    return records


def merge_records(old: list[dict[str, Any]], current: list[dict[str, Any]]) -> list[dict[str, Any]]:
    merged: dict[tuple[str, ...], dict[str, Any]] = {}
    for record in old + current:
        key = record_key(record)
        merged[key] = {**merged.get(key, {}), **{k: v for k, v in record.items() if v not in (None, "")}}
    return list(merged.values())


def assign_ids(records: list[dict[str, Any]], week: str) -> tuple[dict[str, str], dict[tuple[str, ...], str]]:
    counters = {letter: 0 for letter in "CFXO"}
    old_to_new: dict[str, str] = {}
    key_to_new: dict[tuple[str, ...], str] = {}
    for record in records:
        existing = str(record.get("文献编号") or record.get("文献ID") or "")
        category = category_for(record)
        existing_match = re.fullmatch(r"\d{4}-W\d{2}-([CFXO])(\d+)", existing)
        if existing_match:
            new_id = existing
            if existing.startswith(week + "-"):
                letter = existing_match.group(1)
                counters[letter] = max(counters[letter], int(existing_match.group(2)))
        else:
            counters[category] += 1
            new_id = f"{week}-{category}{counters[category]:02d}"
        if existing:
            old_to_new[existing] = new_id
        key_to_new[record_key(record)] = new_id
    return old_to_new, key_to_new


def to_row(record: dict[str, Any], doc_id: str, week: str) -> dict[str, Any]:
    category = category_for(record)
    source_fulltext = str(record.get("全文访问状态") or "")
    local_path = str(record.get("本地全文路径") or "")
    existing_doc_path = str(record.get("精读文档路径") or "")
    level = str(record.get("阅读等级") or "")
    if existing_doc_path or (local_path and (level == "A" or "精读" in str(record.get("阅读状态") or ""))):
        fulltext = source_fulltext or "已保存到项目"
        deep_state = "已完成全文精读"
    else:
        fulltext = "仅摘要，待全文"
        deep_state = "仅摘要，待全文"
    abstract = str(record.get("原文摘要") or "").strip()
    if abstract in {"原文摘要未获取", "未报告", "需全文核验"}:
        abstract = ""
    chinese_abstract = str(record.get("中文摘要") or "").strip()
    if chinese_abstract in {"未报告", "需全文核验"}:
        chinese_abstract = ""
    return {
        "文献编号": doc_id,
        "文献类别": category,
        "文章类型": record.get("文章类型") or record.get("材料类型") or "待核验",
        "英文题名": record.get("英文标题") or record.get("英文题名") or "待核验",
        "中文题名": record.get("中文标题") or record.get("中文题名") or "",
        "期刊": record.get("期刊") or record.get("来源或期刊") or "待核验",
        "第一作者": record.get("第一作者") or str(record.get("作者") or "").split(";")[0] or "待核验",
        "DOI": record.get("DOI") or "待核验",
        "正式发表日期": record.get("正式发表日期") or record.get("Online first日期") or "待核验",
        "原文摘要": abstract,
        "中文摘要": chinese_abstract,
        "全文访问状态": fulltext,
        "精读状态": deep_state,
        "精读文档路径": existing_doc_path,
        "备注": "出版社页面、DOI和文章类型按当前可核验信息记录。",
    }


def update(root: Path, bundle_path: Path, week: str) -> dict[str, Any]:
    workbook = root / "library/master_literature_review.xlsx"
    old = read_old_records(workbook)
    bundle = json.loads(bundle_path.read_text(encoding="utf-8"))
    current = bundle.get("literature_records") or bundle.get("records") or []
    records = merge_records(old, current)
    old_to_new, key_to_new = assign_ids(records, week)
    needs_compact = not workbook.exists()
    preserved_sheets: dict[str, list[list[Any]]] = {}
    if workbook.exists():
        probe = load_workbook(workbook, read_only=True)
        needs_compact = set(probe.sheetnames) != set(SHEET_SCHEMAS)
        if not needs_compact:
            for sheet_name, columns in SHEET_SCHEMAS.items():
                if [cell.value for cell in probe[sheet_name][1]] != columns:
                    needs_compact = True
                    break
        probe.close()
    if needs_compact:
        if workbook.exists():
            source = load_workbook(workbook, read_only=True, data_only=False)
            for sheet_name in ("选题池", "研究趋势", "修改日志"):
                if sheet_name in source.sheetnames:
                    preserved_sheets[sheet_name] = [
                        [cell.value for cell in row]
                        for row in source[sheet_name].iter_rows(values_only=False)
                    ]
            source.close()
            backup_dir = workbook.parent / "backups"
            backup_dir.mkdir(exist_ok=True)
            shutil.copy2(workbook, backup_dir / f"{workbook.stem}_pre_compact_{date.today().isoformat()}.xlsx")
        compact_tmp = workbook.with_name(workbook.stem + "_compact_tmp.xlsx")
        create_compact_workbook(compact_tmp)
        compact_tmp.replace(workbook)
    wb = load_workbook(workbook)
    for sheet_name, matrix in preserved_sheets.items():
        target = wb[sheet_name]
        if target.max_row > 1:
            target.delete_rows(2, target.max_row - 1)
        for values in matrix[1:]:
            if any(value not in (None, "") for value in values):
                target.append(values)
    ws = wb["文献表"]
    ws.delete_rows(2, ws.max_row)
    rows = []
    for record in records:
        row = to_row(record, key_to_new[record_key(record)], week)
        rows.append(row)
        ws.append([row.get(field, "") for field in LITERATURE_COLUMNS])
    log = wb["修改日志"]
    append_change(log, week, "migrated", "工作簿结构", "旧多表结构", "文献表、选题池、研究趋势、修改日志", str(bundle_path), "update_literature_workbook.py", "删除旧字段和旧分类；保留文献与阅读证据")
    for row in rows:
        append_change(log, row["文献编号"], "upsert", "文献表", "", row["英文题名"], str(bundle_path), "update_literature_workbook.py")
    backup = atomic_save(wb, workbook, make_backup=True)
    return {"records": len(rows), "id_map": old_to_new, "backup": str(backup) if backup else None, "categories": {c: sum(r["文献类别"] == c for r in rows) for c in "CFXO"}}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("bundle_json")
    parser.add_argument("--week", default=f"{date.today().isocalendar().year}-W{date.today().isocalendar().week:02d}")
    parser.add_argument("--project-root")
    args = parser.parse_args()
    result = update(project_root(args.project_root), Path(args.bundle_json), args.week)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
