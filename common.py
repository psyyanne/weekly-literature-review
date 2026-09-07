"""Shared compact schemas and safe workbook helpers for weekly literature review."""
from __future__ import annotations

import re
import shutil
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_PROJECT_ROOT = SKILL_DIR.parents[2]
MASTER_NAME = "master_literature_review.xlsx"

LITERATURE_COLUMNS = [
    "文献编号", "文献类别", "文章类型", "英文题名", "中文题名", "期刊", "第一作者", "DOI", "正式发表日期",
    "原文摘要", "中文摘要", "全文访问状态", "精读状态", "精读文档路径", "备注",
]
TOPIC_COLUMNS = [
    "选题编号", "来源文献编号", "备选研究题目", "所属心理学方向", "研究缺口", "理论基础",
    "核心变量", "研究对象", "方法建议", "相比原文的推进", "创新点", "可行性",
    "需要继续核验的问题", "首次提出日期", "最近更新日期", "选题状态", "备注",
]
TREND_COLUMNS = [
    "趋势编号", "趋势名称", "所属心理学方向", "支持文献编号", "主要理论", "主要变量", "常用研究方法",
    "主要结论", "不一致结果", "近期变化", "研究缺口", "可以继续发展的方向", "证据强度", "趋势状态",
    "最近更新日期", "备注",
]
CHANGE_COLUMNS = ["修改时间", "对象编号", "修改类型", "修改字段", "原内容", "新内容", "修改原因", "数据来源", "执行脚本", "备注"]
SHEET_SCHEMAS = {"文献表": LITERATURE_COLUMNS, "选题池": TOPIC_COLUMNS, "研究趋势": TREND_COLUMNS, "修改日志": CHANGE_COLUMNS}
READABLE_FULLTEXT = {"可直接阅读全文", "可直接下载PDF", "用户已提供全文", "已保存到项目"}
TOPIC_STATES = ["新提出", "需要补充文献", "理论框架形成中", "方法设计形成中", "可进一步发展", "已被近期研究覆盖", "暂缓", "放弃"]

def project_root(value: str | Path | None = None) -> Path:
    return Path(value).resolve() if value else DEFAULT_PROJECT_ROOT

def normalize_doi(value: Any) -> str:
    text = str(value or "").strip().lower()
    if text in {"", "待核验", "未报告", "无", "不适用"}:
        return ""
    return re.sub(r"^(https?://(dx\.)?doi\.org/|doi:\s*)", "", text).rstrip(" .;,)")

def normalize_key(value: Any) -> str:
    return re.sub(r"[^\w\u4e00-\u9fff]+", "", str(value or "").lower())

def record_key(record: dict[str, Any]) -> tuple[str, ...]:
    doi = normalize_doi(record.get("DOI"))
    if doi:
        return ("doi", doi)
    title = record.get("英文题名") or record.get("英文标题") or record.get("题名")
    return ("meta", normalize_key(title), normalize_key(record.get("第一作者") or record.get("作者")))

def worksheet_records(ws: Any) -> list[dict[str, Any]]:
    headers = [cell.value for cell in ws[1]]
    return [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True) if any(value not in (None, "") for value in row)]

def upsert_dict(ws: Any, key_name: str, record: dict[str, Any]) -> tuple[str, int]:
    headers = [cell.value for cell in ws[1]]
    key_col = headers.index(key_name) + 1
    key_value = str(record.get(key_name) or "")
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, key_col).value or "") == key_value:
            for col, header in enumerate(headers, 1):
                if record.get(header) not in (None, ""):
                    ws.cell(row, col).value = record[header]
            return "updated", row
    ws.append([record.get(header, "") for header in headers])
    return "added", ws.max_row

def style_sheet(ws: Any) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.protection = Protection(locked=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for column in ws.columns:
        letter = column[0].column_letter
        max_len = max((len(str(c.value or "")) for c in column), default=12)
        ws.column_dimensions[letter].width = min(max(12, max_len + 2), 42)
        for cell in column[1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.protection = Protection(locked=False)
    ws.row_dimensions[1].height = 32

def add_validations(ws: Any, columns: list[str]) -> None:
    if "文献类别" in columns:
        col = columns.index("文献类别") + 1
        dv = DataValidation(type="list", formula1='"C,F,X,O"', allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(f"{ws.cell(1, col).column_letter}2:{ws.cell(1, col).column_letter}1048576")
    if "选题状态" in columns:
        col = columns.index("选题状态") + 1
        dv = DataValidation(type="list", formula1=f'"{",".join(TOPIC_STATES)}"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{ws.cell(1, col).column_letter}2:{ws.cell(1, col).column_letter}1048576")

def atomic_save(wb: Any, path: Path, make_backup: bool = True) -> Path | None:
    path.parent.mkdir(parents=True, exist_ok=True)
    backup = None
    if make_backup and path.exists():
        backup_dir = path.parent / "backups"
        backup_dir.mkdir(exist_ok=True)
        backup = backup_dir / f"{path.stem}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        shutil.copy2(path, backup)
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as handle:
        temp_path = Path(handle.name)
    try:
        wb.save(temp_path)
        check = load_workbook(temp_path, read_only=True)
        check.close()
        temp_path.replace(path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise
    return backup

def create_compact_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "文献表"
    for name, columns in SHEET_SCHEMAS.items():
        target = ws if name == "文献表" else wb.create_sheet(name)
        target.append(columns)
        style_sheet(target)
        add_validations(target, columns)
    atomic_save(wb, path, make_backup=False)

def append_change(ws: Any, object_id: str, change_type: str, field: str, old: Any, new: Any, source: str, script: str, note: str = "") -> None:
    ws.append([datetime.now().isoformat(timespec="seconds"), object_id, change_type, field, old or "", new or "", "持续文献阅读与资料更新", source, script, note])
